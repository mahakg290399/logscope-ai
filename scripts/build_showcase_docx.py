"""Build docs/LogScope-AI-Showcase.docx from docs/showcase screenshots.

Usage:  python scripts/build_showcase_docx.py
Input:  docs/showcase/shot-dashboard.png, docs/showcase/shot-triage.png
Output: docs/LogScope-AI-Showcase.docx (upload-ready portfolio case study)
"""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
SHOWCASE = ROOT / "docs" / "showcase"
OUT = ROOT / "docs" / "LogScope-AI-Showcase.docx"

ACCENT = RGBColor(0x6D, 0x28, 0xD9)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = ACCENT if level == 1 else RGBColor(0x1F, 0x29, 0x37)
    return h


def figure(doc, path, caption):
    doc.add_picture(str(path), width=Inches(6.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)


def main():
    for name in ("shot-dashboard.png", "shot-triage.png"):
        if not (SHOWCASE / name).exists():
            raise SystemExit(f"missing input: docs/showcase/{name}")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("LogScope AI — SRE Log Intelligence Platform", level=0)
    sub = doc.add_paragraph()
    run = sub.add_run("Case study · v2.0 · Live in production at logscope.freengineer.me")
    run.italic = True
    run.font.size = Pt(11)

    doc.add_paragraph(
        "LogScope AI turns raw application logs into triaged incidents. It tails logs from "
        "multiple services, redacts PII and secrets before anything leaves the collector, mines "
        "repetitive lines into Drain3 templates, streams sanitized events through Kafka, aggregates "
        "5-minute windows, and fires deterministic anomaly rules (new templates, z-score frequency "
        "spikes, error bursts). An LLM copilot then judges each incident with structured Sev-0…Sev-4 "
        "verdicts, evidence, and recommended actions on a real-time dashboard."
    )

    heading(doc, "Live dashboard — health at a glance", level=1)
    figure(
        doc,
        SHOWCASE / "shot-dashboard.png",
        "Figure 1 — KPI row (ingestion rate, buffer lag & drops, active anomalies, Drain3 clusters), "
        "5-minute ingestion timeline with error-spike overlay, incident filters, and the Ask SRE Copilot panel.",
    )
    doc.add_paragraph(
        "The KPI row answers four questions instantly: how fast are we ingesting, is anything queued "
        "or dropped, how many incidents need attention, and how many log templates have been mined. "
        "The timeline chart shows observation volume against error spikes per 5-minute bucket, so a "
        "heartbeat arrhythmia is visible before anyone reads a single log line."
    )

    heading(doc, "Incident triage — from signal to action", level=1)
    figure(
        doc,
        SHOWCASE / "shot-triage.png",
        "Figure 2 — Triaged incident cards (INC numbers, severity/confidence badges, template with Drill "
        "Down, causes/actions/evidence), the Ask SRE Copilot chat, and read-only monitored sources.",
    )
    doc.add_paragraph(
        "Every anomaly becomes a tracked incident with a unique INC number, severity (Sev-0 critical "
        "through Sev-4 info) and confidence level kept as separate judgments. Cards carry the mined "
        "template with one-click drill-down into bucket history and sanitized samples, plus diagnostic "
        "causes, recommended actions, retry-triage and a three-outcome resolution flow (Resolved with "
        "fix notes indexed into AI memory, False Alarm, Transient). Nothing is ever deleted — every "
        "verdict stays queryable for future triage."
    )

    heading(doc, "Engineering highlights", level=1)
    for bullet in [
        "Sanitize-first pipeline: raw log lines never reach Kafka or any LLM; redaction and client-ID hashing happen in the collector.",
        "Deterministic detection before AI: math fires the alerts, the model only judges whether investigation is warranted — with prompt-injection guarding.",
        "Bounded everything: capped evidence samples and drop-oldest backpressure with visible counters instead of silent stalls.",
        "Zero-cost production: Terraform-managed Oracle Always Free VM, OCI Vault secrets via instance principal, Caddy HTTPS, GitHub Actions CI/CD.",
        "Verified: 34/34 tests passing, 1,000+ lines/sec sustained throughput with zero drops.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    heading(doc, "Links", level=1)
    doc.add_paragraph("Live dashboard: https://logscope.freengineer.me/")
    doc.add_paragraph("Technical whitepaper: https://logscope.freengineer.me/whitepaper")
    doc.add_paragraph("Source code: https://github.com/mahakg290399/logscope-ai")

    heading(doc, "Tech stack", level=1)
    doc.add_paragraph(
        "Python · FastAPI · Apache Kafka (KRaft) · Drain3 · SQLite (WAL) · Docker · Terraform · "
        "Oracle Cloud Infrastructure · GitHub Actions · OpenAI / NVIDIA NIM / Gemini · Caddy · Chart.js"
    )

    doc.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
